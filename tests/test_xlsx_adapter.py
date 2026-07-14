"""Safe XLSX adapter behavior for DataSource v1."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from data_viewer.domain import (
    DataDomain,
    DataViewerError,
    ErrorCode,
    OperationScope,
    ResourceId,
    SourceCapability,
    StructuredPayload,
    TablePayload,
)
from data_viewer.sources.api import ReadRequest
from data_viewer.sources.registry import SourceRegistry
from data_viewer.sources.xlsx import XLSXAdapter, XLSXSourceSession
from data_viewer.tasks import CancellationToken


def _write_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["name", "value", "formula"])
    sheet.append(["alpha", 2, "=SUM(1, 2)"])
    sheet["E5"] = "lonely"
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "merged"
    sheet.add_table(Table(displayName="DataTable", ref="A1:C2"))
    workbook.defined_names.add(DefinedName("AnswerCell", attr_text="Data!$B$2"))
    unicode_sheet = workbook.create_sheet("Unicode Ω")
    unicode_sheet["A1"] = "snowman ☃"
    workbook.save(path)
    return path


def test_xlsx_adapter_opens_workbook_tree_and_sheet_metadata(tmp_path: Path) -> None:
    path = _write_workbook(tmp_path / "book.xlsx")
    registry = SourceRegistry([XLSXAdapter()])
    session = registry.open(path, cancellation=CancellationToken())
    root = session.root()
    root_page = session.list_children(
        root.resource_id,
        cursor=None,
        page_size=10,
        cancellation=CancellationToken(),
    )
    sheets_page = session.list_children(
        ResourceId(session.source_uri, "/sheets"),
        cursor=None,
        page_size=10,
        cancellation=CancellationToken(),
    )
    data_metadata = session.get_metadata(
        ResourceId(session.source_uri, "/sheets/Data"),
        cancellation=CancellationToken(),
    )
    root_metadata = session.get_metadata(root.resource_id, cancellation=CancellationToken())

    assert root.domain is DataDomain.WORKBOOK
    assert [item.resource_id.node_path for item in root_page.items] == [
        "/defined_names",
        "/sheets",
    ]
    assert [item.resource_id.node_path for item in sheets_page.items] == [
        "/sheets/Data",
        "/sheets/Unicode Ω",
    ]
    assert root_metadata.attributes["external_links_disabled"] is True
    assert root_metadata.attributes["macros_supported"] is False
    assert data_metadata.domain is DataDomain.TABLE
    assert data_metadata.shape == (5, 5)
    assert data_metadata.attributes["dimension"] == "A1:E5"
    assert data_metadata.attributes["tables"] == [{"name": "DataTable", "ref": "A1:C2"}]
    assert data_metadata.attributes["merged_ranges"] == ["A4:B4"]
    assert data_metadata.attributes["formula_cells"] == [
        {
            "coordinate": "C2",
            "formula": "=SUM(1, 2)",
            "cached_value": None,
            "cached_available": False,
        }
    ]
    assert SourceCapability.EDIT_PATCH not in data_metadata.capabilities
    assert SourceCapability.ATOMIC_REWRITE not in data_metadata.capabilities
    session.close()


def test_xlsx_reads_bounded_sheet_pages_and_defined_names(tmp_path: Path) -> None:
    path = _write_workbook(tmp_path / "pages.xlsx")
    session = XLSXAdapter().open(path, cancellation=CancellationToken())
    page_result = session.read(
        ReadRequest(
            resource_id=ResourceId(session.source_uri, "/sheets/Data"),
            scope=OperationScope.PAGE,
            row_offset=1,
            row_limit=2,
            selected_columns=("A", "C", "D"),
        ),
        cancellation=CancellationToken(),
        progress=lambda _done, _total, _message: None,
    )
    name_result = session.read(
        ReadRequest(
            resource_id=ResourceId(session.source_uri, "/defined_names/AnswerCell"),
            scope=OperationScope.FULL,
        ),
        cancellation=CancellationToken(),
        progress=lambda _done, _total, _message: None,
    )

    assert isinstance(page_result.payload, TablePayload)
    assert [column.name for column in page_result.payload.columns] == ["A", "C", "D"]
    assert page_result.payload.row_offset == 1
    assert page_result.payload.total_rows == 5
    assert list(page_result.payload.column_values[0]) == ["alpha", None]
    assert list(page_result.payload.column_values[1]) == ["=SUM(1, 2)", None]
    assert list(page_result.payload.column_values[2]) == [None, None]
    assert isinstance(name_result.payload, StructuredPayload)
    assert name_result.payload.value == {
        "name": "AnswerCell",
        "attr_text": "Data!$B$2",
        "destinations": [{"sheet": "Data", "range": "$B$2"}],
    }
    session.close()


def test_xlsx_exposes_cell_table_and_merge_inspection_nodes(tmp_path: Path) -> None:
    path = _write_workbook(tmp_path / "inspection.xlsx")
    session = XLSXAdapter().open(path, cancellation=CancellationToken())
    sheet_page = session.list_children(
        ResourceId(session.source_uri, "/sheets/Data"),
        cursor=None,
        page_size=10,
        cancellation=CancellationToken(),
    )
    cell_result = session.read(
        ReadRequest(
            resource_id=ResourceId(session.source_uri, "/sheets/Data/cells/C2"),
            scope=OperationScope.FULL,
        ),
        cancellation=CancellationToken(),
        progress=lambda _done, _total, _message: None,
    )
    table_result = session.read(
        ReadRequest(
            resource_id=ResourceId(session.source_uri, "/sheets/Data/tables/DataTable"),
            scope=OperationScope.FULL,
        ),
        cancellation=CancellationToken(),
        progress=lambda _done, _total, _message: None,
    )
    merge_result = session.read(
        ReadRequest(
            resource_id=ResourceId(session.source_uri, "/sheets/Data/merged_ranges/A4:B4"),
            scope=OperationScope.FULL,
        ),
        cancellation=CancellationToken(),
        progress=lambda _done, _total, _message: None,
    )

    assert [item.resource_id.node_path for item in sheet_page.items] == [
        "/sheets/Data/cells",
        "/sheets/Data/merged_ranges",
        "/sheets/Data/tables",
    ]
    assert isinstance(cell_result.payload, StructuredPayload)
    assert cell_result.payload.value == {
        "coordinate": "C2",
        "value": "=SUM(1, 2)",
        "data_type": "formula",
        "is_formula": True,
        "cached_value": None,
        "cached_available": False,
    }
    assert isinstance(table_result.payload, StructuredPayload)
    assert table_result.payload.value == {"name": "DataTable", "ref": "A1:C2"}
    assert isinstance(merge_result.payload, StructuredPayload)
    assert merge_result.payload.value == {"range": "A4:B4", "top_left": "A4"}
    session.close()


def test_xlsx_rejects_xlsm_encrypted_malformed_budget_and_closed_session(
    tmp_path: Path,
) -> None:
    xlsm = tmp_path / "macro.xlsm"
    xlsm.write_bytes(b"PK\x03\x04")
    encrypted = tmp_path / "encrypted.xlsx"
    encrypted.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 64)
    malformed = tmp_path / "bad.xlsx"
    malformed.write_bytes(b"not a zip workbook")
    path = _write_workbook(tmp_path / "budget.xlsx")

    assert XLSXAdapter().probe(xlsm, xlsm.read_bytes()) is None

    with pytest.raises(DataViewerError) as encrypted_error:
        XLSXAdapter().open(encrypted, cancellation=CancellationToken())
    assert encrypted_error.value.code is ErrorCode.SOURCE_ENCRYPTED

    with pytest.raises(DataViewerError) as malformed_error:
        XLSXAdapter().open(malformed, cancellation=CancellationToken())
    assert malformed_error.value.code is ErrorCode.SOURCE_MALFORMED

    with pytest.raises(DataViewerError) as budget_error:
        XLSXSourceSession(path, max_cells=1)
    assert budget_error.value.code is ErrorCode.BUDGET_EXCEEDED

    session = XLSXAdapter().open(path, cancellation=CancellationToken())
    session.close()
    with pytest.raises(DataViewerError) as closed_error:
        session.root()
    assert closed_error.value.code is ErrorCode.SOURCE_CLOSED
