"""Session implementation for read-only safe XLSX workbooks."""

from __future__ import annotations

import json
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import BadZipFile

import numpy as np
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from data_viewer.domain import (
    ColumnSpec,
    DataDomain,
    DataMetadata,
    DataViewerError,
    ErrorCode,
    JsonValue,
    NodeKind,
    OperationScope,
    ReadResult,
    ResourceId,
    ResourceNode,
    SourceCapability,
    SourceFingerprint,
    StructuredPayload,
    TablePayload,
)
from data_viewer.sources.api import NodePage, ProgressCallback, ReadRequest

from .constants import OLE_COMPOUND_SIGNATURE

DEFAULT_MAX_CELLS = 1_000_000
DEFAULT_PAGE_ROWS = 1_000


@dataclass(slots=True)
class _XLSXNode:
    path: str
    name: str
    domain: DataDomain
    node_kind: NodeKind
    kind: str
    attributes: dict[str, JsonValue] = field(default_factory=dict)
    children: dict[str, "_XLSXNode"] = field(default_factory=dict)

    @property
    def has_children(self) -> bool:
        return bool(self.children)


class XLSXSourceSession:
    """One owned, read-only XLSX workbook session."""

    def __init__(
        self,
        path: Path,
        *,
        max_cells: int = DEFAULT_MAX_CELLS,
    ) -> None:
        self._path = path
        self._source_uri = path.resolve(strict=False).as_uri()
        self._closed = False
        self._fingerprint = _fingerprint(path)
        self._max_cells = _positive(max_cells, "max_cells")
        self._workbook: Workbook
        self._cached_workbook: Workbook
        self._open_workbooks()
        self._check_cell_budget()
        self._root = self._build_tree()
        self._nodes = dict(_walk_nodes(self._root))

    @property
    def source_uri(self) -> str:
        return self._source_uri

    @property
    def fingerprint(self) -> SourceFingerprint:
        return self._fingerprint

    @property
    def capabilities(self) -> SourceCapability:
        return SourceCapability.HIERARCHY | SourceCapability.SEARCH | SourceCapability.SAVE_AS

    def root(self) -> ResourceNode:
        self._ensure_open(operation="source.xlsx.root")
        return self._resource_node(self._root)

    def list_children(
        self,
        parent: ResourceId,
        *,
        cursor: str | None,
        page_size: int,
        cancellation: object,
    ) -> NodePage:
        self._ensure_open(operation="source.xlsx.list_children")
        _raise_if_cancelled(cancellation, operation="source.xlsx.list_children")
        if page_size <= 0:
            raise ValueError("page_size must be a positive integer")
        parent_node = self._node_for(parent, operation="source.xlsx.list_children")
        items = tuple(parent_node.children.values())
        start = _cursor_to_index(cursor)
        stop = min(start + page_size, len(items))
        return NodePage(
            items=tuple(self._resource_node(node) for node in items[start:stop]),
            next_cursor=str(stop) if stop < len(items) else None,
            total_count=len(items),
        )

    def get_metadata(
        self,
        resource: ResourceId,
        *,
        cancellation: object,
    ) -> DataMetadata:
        self._ensure_open(operation="source.xlsx.get_metadata")
        _raise_if_cancelled(cancellation, operation="source.xlsx.get_metadata")
        node = self._node_for(resource, operation="source.xlsx.get_metadata")
        attributes: dict[str, JsonValue] = {
            "format": "XLSX",
            "read_only": True,
            "source_fingerprint": self._fingerprint.to_json(),
        }
        attributes.update(node.attributes)
        columns: tuple[ColumnSpec, ...] = ()
        shape: tuple[int, ...] = ()
        dtype = node.kind
        if node.kind == "sheet":
            sheet = self._sheet(node.name)
            shape = (int(sheet.max_row), int(sheet.max_column))
            columns = _columns_for_sheet(sheet)
            dtype = "worksheet"
        elif node.path == "/":
            shape = (len(self._workbook.sheetnames),)
            dtype = "workbook"
        return DataMetadata(
            resource_id=ResourceId(self._source_uri, node.path),
            name=node.name,
            domain=node.domain,
            node_kind=node.node_kind,
            shape=shape,
            dtype=dtype,
            logical_size_bytes=None,
            storage_size_bytes=self._path.stat().st_size,
            capabilities=self._capabilities_for(node),
            attributes=attributes,
            columns=columns,
        )

    def read(
        self,
        request: ReadRequest,
        *,
        cancellation: object,
        progress: ProgressCallback,
    ) -> ReadResult:
        self._ensure_open(operation="source.xlsx.read")
        _raise_if_cancelled(cancellation, operation="source.xlsx.read")
        node = self._node_for(request.resource_id, operation="source.xlsx.read")
        progress(0, 1, "start")
        if node.kind == "sheet":
            result = self._read_sheet(node, request)
        elif node.kind in {"defined_name", "cell", "table", "merged_range"}:
            result = self._read_structured(node, request)
        else:
            raise DataViewerError(
                code=ErrorCode.CAPABILITY_UNAVAILABLE,
                message="XLSX container resources cannot be read as payloads.",
                operation="source.xlsx.read",
                resource_id=request.resource_id,
            )
        progress(1, 1, "done")
        return result

    def search(
        self,
        query: str,
        *,
        cursor: str | None,
        page_size: int,
        cancellation: object,
    ) -> NodePage:
        self._ensure_open(operation="source.xlsx.search")
        _raise_if_cancelled(cancellation, operation="source.xlsx.search")
        if page_size <= 0:
            raise ValueError("page_size must be a positive integer")
        normalized = query.strip().lower()
        if not normalized:
            return NodePage(items=(), next_cursor=None, total_count=0)
        matches = tuple(
            self._resource_node(node)
            for path, node in self._nodes.items()
            if path != "/" and (normalized in path.lower() or normalized in node.name.lower())
        )
        start = _cursor_to_index(cursor)
        stop = min(start + page_size, len(matches))
        return NodePage(
            items=matches[start:stop],
            next_cursor=str(stop) if stop < len(matches) else None,
            total_count=len(matches),
        )

    def refresh_fingerprint(self) -> SourceFingerprint:
        self._ensure_open(operation="source.xlsx.refresh_fingerprint")
        self._fingerprint = _fingerprint(self._path)
        return self._fingerprint

    def close(self) -> None:
        self._closed = True
        self._workbook.close()
        self._cached_workbook.close()

    def _open_workbooks(self) -> None:
        prefix = self._path.read_bytes()[:8]
        if prefix.startswith(OLE_COMPOUND_SIGNATURE):
            raise DataViewerError(
                code=ErrorCode.SOURCE_ENCRYPTED,
                message="Encrypted or legacy OLE workbooks are not supported as XLSX.",
                operation="source.xlsx.open",
                details={"path": str(self._path), "encrypted_or_ole": True},
            )
        try:
            self._workbook = load_workbook(
                self._path,
                read_only=False,
                data_only=False,
                keep_links=False,
            )
            self._cached_workbook = load_workbook(
                self._path,
                read_only=False,
                data_only=True,
                keep_links=False,
            )
        except BadZipFile:
            raise
        except Exception as exc:
            raise DataViewerError(
                code=ErrorCode.SOURCE_MALFORMED,
                message="XLSX source is malformed or unsupported.",
                operation="source.xlsx.open",
                details={"path": str(self._path)},
                cause=exc,
            ) from exc

    def _check_cell_budget(self) -> None:
        total = 0
        for sheet in self._workbook.worksheets:
            total += int(sheet.max_row) * int(sheet.max_column)
            if total > self._max_cells:
                raise DataViewerError(
                    code=ErrorCode.BUDGET_EXCEEDED,
                    message="XLSX workbook cell dimensions exceed the configured budget.",
                    operation="source.xlsx.validate_budget",
                    details={"cell_count": total, "max_cells": self._max_cells},
                )

    def _build_tree(self) -> _XLSXNode:
        root = _XLSXNode(
            path="/",
            name=self._path.name,
            domain=DataDomain.WORKBOOK,
            node_kind=NodeKind.ROOT,
            kind="workbook",
            attributes={
                "sheet_count": len(self._workbook.sheetnames),
                "external_links_disabled": True,
                "macros_supported": False,
            },
        )
        sheets_node = _XLSXNode(
            path="/sheets",
            name="sheets",
            domain=DataDomain.WORKBOOK,
            node_kind=NodeKind.CONTAINER,
            kind="sheets",
        )
        defined_names_node = _XLSXNode(
            path="/defined_names",
            name="defined_names",
            domain=DataDomain.METADATA,
            node_kind=NodeKind.CONTAINER,
            kind="defined_names",
        )
        root.children["defined_names"] = defined_names_node
        root.children["sheets"] = sheets_node
        for sheet_name in self._workbook.sheetnames:
            sheet = self._sheet(sheet_name)
            sheet_node = self._sheet_node(sheet)
            sheets_node.children[sheet_name] = sheet_node
        for name in sorted(self._workbook.defined_names.keys()):
            node = _XLSXNode(
                path=_pointer_join("/defined_names", name),
                name=name,
                domain=DataDomain.STRUCTURED,
                node_kind=NodeKind.RESOURCE,
                kind="defined_name",
                attributes=_defined_name_payload(self._workbook.defined_names[name]),
            )
            defined_names_node.children[name] = node
        return root

    def _sheet_node(self, sheet: Worksheet) -> _XLSXNode:
        path = _pointer_join("/sheets", sheet.title)
        node = _XLSXNode(
            path=path,
            name=sheet.title,
            domain=DataDomain.TABLE,
            node_kind=NodeKind.RESOURCE,
            kind="sheet",
            attributes=_sheet_attributes(sheet, self._cached_sheet(sheet.title)),
        )
        cells = _XLSXNode(
            path=_pointer_join(path, "cells"),
            name="cells",
            domain=DataDomain.METADATA,
            node_kind=NodeKind.CONTAINER,
            kind="cells",
        )
        tables = _XLSXNode(
            path=_pointer_join(path, "tables"),
            name="tables",
            domain=DataDomain.METADATA,
            node_kind=NodeKind.CONTAINER,
            kind="tables",
        )
        merged_ranges = _XLSXNode(
            path=_pointer_join(path, "merged_ranges"),
            name="merged_ranges",
            domain=DataDomain.METADATA,
            node_kind=NodeKind.CONTAINER,
            kind="merged_ranges",
        )
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and cell.coordinate not in sheet.merged_cells:
                    continue
                cells.children[cell.coordinate] = _XLSXNode(
                    path=_pointer_join(cells.path, cell.coordinate),
                    name=cell.coordinate,
                    domain=DataDomain.STRUCTURED,
                    node_kind=NodeKind.RESOURCE,
                    kind="cell",
                    attributes=_cell_payload(cell, self._cached_sheet(sheet.title)[cell.coordinate]),
                )
        for table in sheet.tables.values():
            tables.children[table.name] = _XLSXNode(
                path=_pointer_join(tables.path, table.name),
                name=table.name,
                domain=DataDomain.STRUCTURED,
                node_kind=NodeKind.RESOURCE,
                kind="table",
                attributes={"name": table.name, "ref": table.ref},
            )
        for merged_range in sorted(str(item) for item in sheet.merged_cells.ranges):
            merged_ranges.children[merged_range] = _XLSXNode(
                path=_pointer_join(merged_ranges.path, merged_range),
                name=merged_range,
                domain=DataDomain.STRUCTURED,
                node_kind=NodeKind.RESOURCE,
                kind="merged_range",
                attributes={
                    "range": merged_range,
                    "top_left": merged_range.split(":", 1)[0],
                },
            )
        node.children["cells"] = cells
        node.children["merged_ranges"] = merged_ranges
        node.children["tables"] = tables
        return node

    def _sheet(self, name: str) -> Worksheet:
        return self._workbook[name]

    def _cached_sheet(self, name: str) -> Worksheet:
        return self._cached_workbook[name]

    def _node_for(self, resource: ResourceId, *, operation: str) -> _XLSXNode:
        if resource.source_uri != self._source_uri:
            raise _resource_not_found(resource, operation=operation)
        node = self._nodes.get(resource.node_path or "/")
        if node is None:
            raise _resource_not_found(resource, operation=operation)
        return node

    def _resource_node(self, node: _XLSXNode) -> ResourceNode:
        return ResourceNode(
            resource_id=ResourceId(self._source_uri, node.path),
            name=node.name,
            node_kind=node.node_kind,
            domain=node.domain,
            has_children=node.has_children,
            summary=f"XLSX {node.kind}",
        )

    def _capabilities_for(self, node: _XLSXNode) -> SourceCapability:
        capabilities = SourceCapability.SEARCH | SourceCapability.SAVE_AS
        if node.has_children or node.node_kind in {NodeKind.ROOT, NodeKind.CONTAINER}:
            capabilities |= SourceCapability.HIERARCHY
        if node.kind == "sheet":
            capabilities |= SourceCapability.PAGED_ROWS | SourceCapability.COLUMN_SCHEMA
        return capabilities

    def _read_sheet(self, node: _XLSXNode, request: ReadRequest) -> ReadResult:
        if request.selection.axes:
            raise _unsupported_selection(request, operation="source.xlsx.read")
        if request.scope not in {OperationScope.PAGE, OperationScope.FULL}:
            raise _unsupported_scope(request, operation="source.xlsx.read")
        sheet = self._sheet(node.name)
        row_offset = request.row_offset or 0
        row_limit = request.row_limit
        if request.scope is OperationScope.PAGE and row_limit is None:
            row_limit = DEFAULT_PAGE_ROWS
        if request.scope is OperationScope.FULL and row_limit is None:
            row_limit = int(sheet.max_row)
        if row_limit is None:
            raise DataViewerError(
                code=ErrorCode.INTERNAL_ERROR,
                message="XLSX row limit was not resolved for the requested read.",
                operation="source.xlsx.read",
                resource_id=request.resource_id,
            )
        selected_columns = request.selected_columns or tuple(
            get_column_letter(index) for index in range(1, int(sheet.max_column) + 1)
        )
        selected_indices = tuple(_column_index(column) for column in selected_columns)
        for column_index in selected_indices:
            if column_index < 1 or column_index > int(sheet.max_column):
                raise DataViewerError(
                    code=ErrorCode.SELECTION_INVALID,
                    message="Selected XLSX column is outside the sheet dimensions.",
                    operation="source.xlsx.read",
                    resource_id=request.resource_id,
                    details={"column_index": column_index, "max_column": int(sheet.max_column)},
                )
        start_row = row_offset + 1
        end_row = min(row_offset + row_limit, int(sheet.max_row))
        row_values: list[tuple[object, ...]] = []
        for row_number in range(start_row, end_row + 1):
            row_values.append(
                tuple(
                    _display_cell_value(sheet.cell(row=row_number, column=column_index))
                    for column_index in selected_indices
                )
            )
        schema = tuple(
            ColumnSpec(
                name=column,
                dtype="object",
                nullable=True,
                metadata={"excel_column": column},
            )
            for column in selected_columns
        )
        payload = TablePayload(
            columns=schema,
            column_values=_column_arrays(schema, row_values),
            row_offset=row_offset,
            total_rows=int(sheet.max_row),
        )
        bytes_read = _table_payload_bytes(payload)
        if bytes_read > request.max_bytes:
            raise DataViewerError(
                code=ErrorCode.BUDGET_EXCEEDED,
                message="XLSX sheet page exceeds the requested read budget.",
                operation="source.xlsx.read",
                resource_id=request.resource_id,
                details={"bytes_read": bytes_read, "max_bytes": request.max_bytes},
            )
        return ReadResult(
            payload=payload,
            scope=request.scope,
            bytes_read=bytes_read,
            is_sampled=False,
            sample=request.sample,
        )

    def _read_structured(self, node: _XLSXNode, request: ReadRequest) -> ReadResult:
        if request.scope is not OperationScope.FULL:
            raise _unsupported_scope(request, operation="source.xlsx.read")
        if request.selection.axes or request.row_offset is not None or request.row_limit is not None:
            raise _unsupported_selection(request, operation="source.xlsx.read")
        value = dict(node.attributes)
        size = len(json.dumps(value, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
        if size > request.max_bytes:
            raise DataViewerError(
                code=ErrorCode.BUDGET_EXCEEDED,
                message="XLSX structured resource exceeds the requested read budget.",
                operation="source.xlsx.read",
                resource_id=request.resource_id,
                details={"bytes_read": size, "max_bytes": request.max_bytes},
            )
        return ReadResult(
            payload=StructuredPayload(value),
            scope=OperationScope.FULL,
            bytes_read=size,
            is_sampled=False,
            sample=request.sample,
        )

    def _ensure_open(self, operation: str) -> None:
        if self._closed:
            raise DataViewerError(
                code=ErrorCode.SOURCE_CLOSED,
                message="XLSX source session is closed.",
                operation=operation,
                details={"source_uri": self._source_uri},
            )


def _sheet_attributes(sheet: Worksheet, cached_sheet: Worksheet) -> dict[str, JsonValue]:
    return {
        "title": sheet.title,
        "dimension": sheet.calculate_dimension(),
        "max_row": int(sheet.max_row),
        "max_column": int(sheet.max_column),
        "merged_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
        "tables": [
            {"name": table.name, "ref": table.ref}
            for table in sorted(sheet.tables.values(), key=lambda item: item.name)
        ],
        "formula_cells": [
            _formula_cell_payload(cell, cached_sheet[cell.coordinate])
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ],
        "cell_type_counts": _cell_type_counts(sheet),
    }


def _cell_payload(cell: Cell, cached_cell: Cell) -> dict[str, JsonValue]:
    payload: dict[str, JsonValue] = {
        "coordinate": cell.coordinate,
        "value": _json_safe(_display_cell_value(cell)),
        "data_type": _cell_data_type(cell),
        "is_formula": cell.data_type == "f",
    }
    if cell.data_type == "f":
        payload["cached_value"] = _json_safe(cached_cell.value)
        payload["cached_available"] = cached_cell.value is not None
    return payload


def _formula_cell_payload(cell: Cell, cached_cell: Cell) -> dict[str, JsonValue]:
    return {
        "coordinate": cell.coordinate,
        "formula": str(cell.value),
        "cached_value": _json_safe(cached_cell.value),
        "cached_available": cached_cell.value is not None,
    }


def _defined_name_payload(defined_name: object) -> dict[str, JsonValue]:
    destinations: list[JsonValue] = []
    try:
        raw_destinations = list(defined_name.destinations)  # type: ignore[attr-defined]
    except Exception:
        raw_destinations = []
    for sheet, cell_range in raw_destinations:
        destinations.append({"sheet": str(sheet), "range": str(cell_range)})
    return {
        "name": str(getattr(defined_name, "name", "")),
        "attr_text": str(getattr(defined_name, "attr_text", "")),
        "destinations": destinations,
    }


def _columns_for_sheet(sheet: Worksheet) -> tuple[ColumnSpec, ...]:
    return tuple(
        ColumnSpec(
            name=get_column_letter(index),
            dtype="object",
            nullable=True,
            metadata={"excel_column": get_column_letter(index)},
        )
        for index in range(1, int(sheet.max_column) + 1)
    )


def _cell_type_counts(sheet: Worksheet) -> dict[str, JsonValue]:
    counts: dict[str, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            name = _cell_data_type(cell)
            counts[name] = counts.get(name, 0) + 1
    return dict(sorted(counts.items()))


def _cell_data_type(cell: Cell) -> str:
    if cell.data_type == "f":
        return "formula"
    if cell.value is None:
        return "blank"
    if cell.data_type == "s":
        return "string"
    if cell.data_type == "n":
        return "number"
    if cell.data_type == "b":
        return "boolean"
    return str(cell.data_type)


def _display_cell_value(cell: Cell) -> object:
    if cell.data_type == "f":
        return cell.value
    return cell.value


def _column_arrays(
    schema: tuple[ColumnSpec, ...],
    rows: list[tuple[object, ...]],
) -> tuple[np.ndarray, ...]:
    return tuple(
        np.array([row[index] for row in rows], dtype=object)
        for index in range(len(schema))
    )


def _table_payload_bytes(payload: TablePayload) -> int:
    total = 0
    for values in payload.column_values:
        total += int(values.nbytes)
        for item in values:
            if item is not None:
                total += len(str(item).encode("utf-8"))
    return total


def _column_index(column: str) -> int:
    column = column.upper()
    index = 0
    for char in column:
        if not "A" <= char <= "Z":
            raise ValueError("selected XLSX columns must use Excel column letters")
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


def _pointer_join(parent: str, token: str) -> str:
    escaped = token.replace("~", "~0").replace("/", "~1")
    if parent in {"", "/"}:
        return "/" + escaped
    return parent.rstrip("/") + "/" + escaped


def _pointer_unescape(token: str) -> str:
    return token.replace("~1", "/").replace("~0", "~")


def _walk_nodes(root: _XLSXNode) -> Iterator[tuple[str, _XLSXNode]]:
    yield root.path, root
    for child in root.children.values():
        yield from _walk_nodes(child)


def _cursor_to_index(cursor: str | None) -> int:
    try:
        start = int(cursor or 0)
    except ValueError as exc:
        raise ValueError("cursor must be a non-negative integer string") from exc
    if start < 0:
        raise ValueError("cursor must be a non-negative integer")
    return start


def _json_safe(value: object) -> JsonValue:
    if isinstance(value, np.generic):
        return _json_safe(value.item())
    if value is None or isinstance(value, str | int | float | bool):
        return value
    return str(value)


def _positive(value: int, label: str) -> int:
    if isinstance(value, bool) or value <= 0:
        raise ValueError(f"{label} must be a positive integer")
    return int(value)


def _unsupported_scope(request: ReadRequest, *, operation: str) -> DataViewerError:
    return DataViewerError(
        code=ErrorCode.CAPABILITY_UNAVAILABLE,
        message="Unsupported XLSX read scope.",
        operation=operation,
        resource_id=request.resource_id,
        details={"scope": request.scope.value},
    )


def _unsupported_selection(request: ReadRequest, *, operation: str) -> DataViewerError:
    return DataViewerError(
        code=ErrorCode.CAPABILITY_UNAVAILABLE,
        message="This XLSX resource does not support the requested selection mode.",
        operation=operation,
        resource_id=request.resource_id,
    )


def _resource_not_found(resource_id: ResourceId, *, operation: str) -> DataViewerError:
    return DataViewerError(
        code=ErrorCode.RESOURCE_NOT_FOUND,
        message="Requested XLSX resource was not found.",
        operation=operation,
        resource_id=resource_id,
        details={"resource_id": resource_id.to_json()},
    )


def _raise_if_cancelled(cancellation: object, *, operation: str) -> None:
    if getattr(cancellation, "is_cancelled", False):
        raise DataViewerError(
            code=ErrorCode.READ_CANCELLED,
            message="XLSX source operation was cancelled.",
            operation=operation,
            retryable=True,
        )


def _fingerprint(path: Path) -> SourceFingerprint:
    try:
        stat = path.stat()
    except OSError as exc:
        raise DataViewerError(
            code=ErrorCode.SOURCE_OPEN_FAILED,
            message="Unable to stat XLSX source.",
            operation="source.xlsx.fingerprint",
            details={"path": str(path)},
            cause=exc,
        ) from exc
    return SourceFingerprint(size_bytes=stat.st_size, modified_time_ns=stat.st_mtime_ns)


__all__ = ["XLSXSourceSession"]
